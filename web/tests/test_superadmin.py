"""Tes fitur superadmin: lihat NIK, ubah, hapus/pulihkan, unduh, kontingen, pengaturan, cadangan."""

from __future__ import annotations

import io

import openpyxl
import pytest
from django.test import Client

from apps.biodata import nik as nik_mod
from apps.biodata.models import JenisPenanda, PenandaPeserta, Peserta
from apps.inti import cadangan, pengaturan
from apps.inti.models import LogAktivitas
from apps.kontingen.models import Kontingen

from .conftest import NIK_CONTOH, data_valid, token_lama


@pytest.fixture
def peserta(client, kontingen_madya):
    client.post("/daftar/", data_valid(kontingen_madya))
    client.cookies.clear()
    return Peserta.objects.get()


def data_ubah(p: Peserta, **ubah) -> dict:
    data = {
        "kontingen": str(p.kontingen_id),
        "tingkat": p.tingkat,
        "peran": p.peran,
        "nama_lengkap": p.nama_lengkap,
        "nik": nik_mod.dekripsi(p.nik_terenkripsi),
        "nomor_induk_anggota": "",
        "jenis_kelamin": p.jenis_kelamin,
        "tempat_lahir": p.tempat_lahir,
        "tanggal_lahir_0": str(p.tanggal_lahir.day),
        "tanggal_lahir_1": str(p.tanggal_lahir.month),
        "tanggal_lahir_2": str(p.tanggal_lahir.year),
        "alamat": p.alamat,
        "golongan_darah": p.golongan_darah,
        "telepon": "",
        "nama_orang_tua": p.nama_orang_tua,
        "telepon_orang_tua": p.telepon_orang_tua,
    }
    data.update(ubah)
    return data


def test_daftar_menyamarkan_nik(klien_superadmin, peserta):
    isi = klien_superadmin.get("/superadmin/biodata/").content.decode()
    assert "6501********0003" in isi and NIK_CONTOH not in isi


def test_cari_4_digit_nik_dan_kode(klien_superadmin, peserta):
    r = klien_superadmin.get("/superadmin/biodata/?q=0003")
    assert r.context["jumlah"] == 1
    r = klien_superadmin.get(f"/superadmin/biodata/?q={peserta.kode_pendaftaran[3:]}")
    assert r.context["jumlah"] == 1
    r = klien_superadmin.get("/superadmin/biodata/?q=9999")
    assert r.context["jumlah"] == 0


def test_lihat_nik_tercatat_di_log(klien_superadmin, peserta):
    r = klien_superadmin.post(f"/superadmin/biodata/{peserta.pk}/nik/")
    assert r.json()["nik"] == NIK_CONTOH
    log = LogAktivitas.objects.get(aksi="nik.dilihat")
    assert log.nama_pelaku == "Penguji"
    assert NIK_CONTOH not in str(log.keterangan)


def test_foto_hanya_lewat_superadmin(klien_superadmin, peserta):
    r = klien_superadmin.get(f"/superadmin/biodata/{peserta.pk}/foto/")
    assert r.status_code == 200 and r["Content-Type"] == "image/jpeg"
    assert "no-store" in r["Cache-Control"]


def test_ubah_biodata_dan_riwayat(klien_superadmin, peserta, kontingen_wira):
    r = klien_superadmin.post(
        f"/superadmin/biodata/{peserta.pk}/ubah/",
        data_ubah(peserta, kontingen=str(kontingen_wira.pk), tempat_lahir="Bulungan"),
    )
    assert r.status_code == 302
    peserta.refresh_from_db()
    assert peserta.kontingen == kontingen_wira and peserta.tempat_lahir == "Bulungan"
    # Pindah ke kontingen Wira tapi tingkat Madya -> tanda tingkat beda.
    assert peserta.penanda.filter(jenis=JenisPenanda.TINGKAT_BEDA).exists()
    log = LogAktivitas.objects.get(aksi="biodata.diubah")
    assert log.data_lama["tempat_lahir"] == "Tanjung Selor" and log.data_baru["tempat_lahir"] == "Bulungan"
    assert NIK_CONTOH not in str(log.data_baru) + str(log.data_lama)


def test_ubah_nik_ke_nik_orang_lain_ditolak(klien_superadmin, client, peserta, kontingen_madya):
    client.post(
        "/daftar/",
        data_valid(
            kontingen_madya, nik="6501014201120004", tanggal_lahir_0="2", nama_lengkap="Rina Dewi", token=token_lama()
        ),
    )
    lain = Peserta.objects.get(nama_lengkap="RINA DEWI")
    r = klien_superadmin.post(f"/superadmin/biodata/{lain.pk}/ubah/", data_ubah(lain, nik=NIK_CONTOH))
    assert r.status_code == 200 and "nik" in r.context["form"].errors


def test_tambah_manual(klien_superadmin, kontingen_madya):
    data = data_valid(kontingen_madya, nama_lengkap="dewi lestari")
    for k in ("token", "persetujuan", "situs_web"):
        data.pop(k)
    r = klien_superadmin.post("/superadmin/biodata/tambah/", data)
    assert r.status_code == 302
    p = Peserta.objects.get()
    assert p.sumber == "SUPERADMIN"
    assert p.nama_lengkap == "DEWI LESTARI"
    assert p.foto_path
    assert LogAktivitas.objects.filter(aksi="biodata.ditambah_superadmin").exists()


def test_tambah_manual_tanpa_foto_ditolak(klien_superadmin, kontingen_madya):
    data = data_valid(kontingen_madya, foto=None)
    for k in ("token", "persetujuan", "situs_web"):
        data.pop(k)
    r = klien_superadmin.post("/superadmin/biodata/tambah/", data)
    assert r.status_code == 200 and "foto" in r.context["form"].errors


def test_ubah_tanpa_foto_baru_memakai_foto_lama(klien_superadmin, peserta):
    r = klien_superadmin.post(
        f"/superadmin/biodata/{peserta.pk}/ubah/", data_ubah(peserta, nama_lengkap="siti aminah putri")
    )
    assert r.status_code == 302
    peserta.refresh_from_db()
    assert peserta.nama_lengkap == "SITI AMINAH PUTRI" and peserta.foto_path


def test_form_superadmin_sama_dengan_form_publik(klien_superadmin, client, kontingen_madya):
    """Tambah manual memakai isian yang sama dengan /daftar/ (cari sekolah, penjelasan peran, pemotong foto)."""
    publik = client.get("/daftar/").content.decode()
    admin = klien_superadmin.get("/superadmin/biodata/tambah/").content.decode()
    for teks in (
        'data-isian="tingkat"', "Cari sekolah", "Guru/pembina yang memimpin kontingen", 'id="data-kegiatan"',
        "data-wadah-potong", "Contoh: Jl. Sengkawit", "Untuk kontak darurat", 'name="foto_potong"', "input uppercase",
    ):  # fmt: skip
        assert teks in publik and teks in admin, teks
    # Yang sengaja berbeda: persetujuan dan anti-bot hanya publik; superadmin boleh sekolah semua tingkat.
    assert 'name="persetujuan"' not in admin and 'name="situs_web"' not in admin
    assert "Tampilkan sekolah semua tingkat" in admin and "Tampilkan sekolah semua tingkat" not in publik


def test_ubah_menampilkan_foto_lama_dan_foto_tidak_wajib(klien_superadmin, peserta):
    isi = klien_superadmin.get(f"/superadmin/biodata/{peserta.pk}/ubah/").content.decode()
    assert f"/superadmin/biodata/{peserta.pk}/foto/" in isi and "Sudah ada foto" in isi
    baris = next(b for b in isi.splitlines() if 'name="foto"' in b)
    assert " required" not in baris


def test_tambah_manual_dengan_potongan_foto(klien_superadmin, kontingen_madya):
    data = data_valid(kontingen_madya, foto_potong='{"x": 100, "y": 50, "w": 300, "h": 400, "r": 0}')
    for k in ("token", "persetujuan", "situs_web"):
        data.pop(k)
    assert klien_superadmin.post("/superadmin/biodata/tambah/", data).status_code == 302
    assert Peserta.objects.get().foto_path


def test_hapus_pulihkan_dan_bentrok_nik(klien_superadmin, client, peserta, kontingen_madya):
    klien_superadmin.post(f"/superadmin/biodata/{peserta.pk}/hapus/")
    assert Peserta.objects.count() == 0 and Peserta.semua.count() == 1
    # Orang lain mendaftar dengan NIK yang sama saat yang lama di tempat sampah.
    client.post("/daftar/", data_valid(kontingen_madya, token=token_lama()))
    r = klien_superadmin.post(f"/superadmin/sampah/biodata/{peserta.pk}/pulihkan/", follow=True)
    assert "sudah dipakai" in r.content.decode()
    assert Peserta.semua.get(pk=peserta.pk).terhapus
    # Setelah yang baru dihapus, yang lama bisa dipulihkan.
    baru = Peserta.objects.get()
    klien_superadmin.post(f"/superadmin/biodata/{baru.pk}/hapus/")
    klien_superadmin.post(f"/superadmin/sampah/biodata/{peserta.pk}/pulihkan/")
    assert not Peserta.semua.get(pk=peserta.pk).terhapus


def test_hapus_permanen_butuh_konfirmasi_kode(klien_superadmin, peserta):
    klien_superadmin.post(f"/superadmin/biodata/{peserta.pk}/hapus/")
    klien_superadmin.post(f"/superadmin/sampah/biodata/{peserta.pk}/hapus-permanen/", {"konfirmasi": "salah"})
    assert Peserta.semua.filter(pk=peserta.pk).exists()
    klien_superadmin.post(
        f"/superadmin/sampah/biodata/{peserta.pk}/hapus-permanen/", {"konfirmasi": peserta.kode_pendaftaran.lower()}
    )
    assert not Peserta.semua.filter(pk=peserta.pk).exists()
    assert LogAktivitas.objects.filter(aksi="biodata.dihapus_permanen").exists()


def test_tanda_ditandai_sudah_dicek(klien_superadmin, client, kontingen_madya):
    # Lahir 2004 = 22 tahun, bukan usia peserta Madya -> tanda usia tidak wajar.
    client.post("/daftar/", data_valid(kontingen_madya, tanggal_lahir_2="2004"))
    t = PenandaPeserta.objects.get(jenis=JenisPenanda.USIA_TIDAK_WAJAR)
    PenandaPeserta.objects.exclude(pk=t.pk).delete()
    klien_superadmin.post(f"/superadmin/tanda/{t.pk}/abaikan/")
    t.refresh_from_db()
    assert t.diabaikan_pada and t.diabaikan_oleh == "Penguji"
    assert klien_superadmin.get("/superadmin/biodata/?tanda=ada").context["jumlah"] == 0


def test_excel_umum_tanpa_data_rahasia(klien_superadmin, peserta):
    r = klien_superadmin.get("/superadmin/biodata/unduh/umum/")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    teks = str([list(row) for row in wb["Semua"].iter_rows(values_only=True)])
    assert "SITI AMINAH" in teks and "6501********0003" in teks
    assert NIK_CONTOH not in teks and "Sengkawit" not in teks and "BUDI SANTOSO" not in teks
    assert set(wb.sheetnames) >= {"Semua", "Mula", "Madya", "Wira", "Kontingen"}
    assert LogAktivitas.objects.filter(aksi="ekspor.umum").exists()


def test_excel_lengkap_wajib_sandi(klien_superadmin, peserta):
    r = klien_superadmin.post("/superadmin/biodata/unduh/lengkap/", {"password": "salah"})
    assert r.status_code == 200 and r.context["form"].errors
    assert LogAktivitas.objects.filter(aksi="ekspor.lengkap_gagal").exists()
    from .conftest import SANDI

    r = klien_superadmin.post("/superadmin/biodata/unduh/lengkap/", {"password": SANDI})
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    teks = str([list(row) for row in wb["Semua"].iter_rows(values_only=True)])
    assert NIK_CONTOH in teks and "Sengkawit" in teks and "BUDI SANTOSO" in teks
    assert LogAktivitas.objects.filter(aksi="ekspor.lengkap").exists()


def test_zip_foto(klien_superadmin, peserta):
    import zipfile

    r = klien_superadmin.get(f"/superadmin/biodata/unduh/foto/?kontingen={peserta.kontingen_id}")
    z = zipfile.ZipFile(io.BytesIO(r.content))
    assert z.namelist() == [f"SITI AMINAH - {peserta.kode_pendaftaran}.jpg"]


def test_ubah_kuota_tercatat_dan_tanda_dihitung_ulang(klien_superadmin, client, kontingen_madya):
    for i, nik in enumerate(["6501014101120003", "6501014201120004"]):
        client.post(
            "/daftar/",
            data_valid(
                kontingen_madya,
                nik=nik,
                tanggal_lahir_0=str(i + 1),
                nama_lengkap=["Ani Lestari", "Bambang Joko"][i],
                token=token_lama(),
            ),
        )
    data = {"kode": "MA-01", "nama": kontingen_madya.nama, "tingkat": "MADYA", "kuota": "1", "catatan": ""}
    klien_superadmin.post(f"/superadmin/kontingen/{kontingen_madya.pk}/ubah/", data)
    assert PenandaPeserta.objects.filter(jenis=JenisPenanda.KUOTA_PENUH).count() == 1
    log = LogAktivitas.objects.get(aksi="kontingen.kuota_diubah")
    assert log.data_lama == {"kuota": 3} and log.data_baru == {"kuota": 1}
    assert Kontingen.objects.get().kuota_diubah_manual


def test_kontingen_berisi_biodata_tidak_bisa_dihapus(klien_superadmin, peserta):
    klien_superadmin.post(f"/superadmin/kontingen/{peserta.kontingen_id}/hapus/")
    assert Kontingen.objects.filter(pk=peserta.kontingen_id).exists()


def test_tambah_kontingen_muncul_di_form(klien_superadmin, client):
    klien_superadmin.post(
        "/superadmin/kontingen/tambah/",
        {"kode": "ma-18", "nama": "SMP Susulan", "tingkat": "MADYA", "kuota": "20", "catatan": ""},
    )
    k = Kontingen.objects.get(kode="MA-18")
    assert "SMP Susulan" in client.get("/daftar/").content.decode()
    assert k.kuota == 20


def test_pengaturan_jadwal_dan_paksa(klien_superadmin, client):
    data = {
        "paksa_status": "TUTUP",
        "buka_pada": "2026-10-26T00:00",
        "tutup_pada": "2026-11-02T23:59",
        "foto_wajib": "on",
        "kontak_nama": "Kak Farhan",
        "kontak_wa": "0812 3456 7890",
        "pengumuman": "",
        "waktu_minimal_detik": "10",
        "tanggal_acuan_usia": "2026-11-12",
    }
    r = klien_superadmin.post("/superadmin/pengaturan/", data)
    assert r.status_code == 302
    assert pengaturan.ambil("kontak.superadmin_wa") == "6281234567890"
    assert pengaturan.ambil("biodata.buka_pada").startswith("2026-10-26T00:00:00+08:00")
    assert client.get("/daftar/").templates[0].name == "biodata/tutup.html"
    assert "wa.me/6281234567890" in client.get("/").content.decode()
    assert LogAktivitas.objects.filter(aksi="pengaturan.diubah").exists()


def test_pengaturan_tutup_sebelum_buka_ditolak(klien_superadmin):
    data = {"paksa_status": "OTOMATIS", "buka_pada": "2026-11-02T00:00", "tutup_pada": "2026-10-26T00:00",
            "kontak_nama": "X", "waktu_minimal_detik": "10", "tanggal_acuan_usia": "2026-11-12"}  # fmt: skip
    r = klien_superadmin.post("/superadmin/pengaturan/", data)
    assert r.status_code == 200 and "tutup_pada" in r.context["form"].errors


def test_mode_perawatan(klien_superadmin):
    pengaturan.simpan("sistem.perawatan", True)
    assert Client().get("/").status_code == 503
    assert klien_superadmin.get("/superadmin/").status_code == 200


def test_cadangan_path_aman():
    assert cadangan.path_aman("../.env") is None
    assert cadangan.path_aman("sigap_..\\..\\x.zip") is None
    assert cadangan.path_aman("bukan.zip") is None
